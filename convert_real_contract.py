"""One-off script: convert a raw SEC EDGAR exhibit HTML into a clean .docx
so it can be run through the app exactly like any other uploaded contract.
Not part of the app itself — just a data-prep utility."""

import re
import sys

from bs4 import BeautifulSoup
from docx import Document


def html_to_docx(html_path, docx_path, title=None):
    with open(html_path, encoding="utf-8", errors="ignore") as f:
        raw = f.read()

    soup = BeautifulSoup(raw, "html.parser")

    # Concatenate text within each <P> block without inserting line breaks —
    # inline tags (e.g. <SUP>th</SUP> in "15th") must stay joined to their
    # surrounding text, not get split onto their own line.
    paragraphs = soup.find_all("p") or [soup]
    lines = []
    for p in paragraphs:
        line = re.sub(r"\s+", " ", p.get_text("")).strip()
        if line:
            lines.append(line)

    doc = Document()
    if title:
        doc.add_heading(title, level=1)
    for line in lines:
        doc.add_paragraph(line)
    doc.save(docx_path)
    print(f"Wrote {docx_path} ({len(lines)} paragraphs)")


if __name__ == "__main__":
    html_to_docx(
        "/tmp/real_msa_raw.html",
        "sample_contracts/real_msa_aspira_h4d.docx",
        title="MASTER SERVICES AGREEMENT — Aspira Women's Health, Inc. / H4D Consulting (SEC EDGAR EX-10.1, filed 2025-09-03)",
    )
