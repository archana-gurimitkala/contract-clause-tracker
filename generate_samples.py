"""One-off script to generate synthetic sample contracts for demo purposes.
These are fabricated documents, not real client contracts."""

from pathlib import Path

from docx import Document
from fpdf import FPDF

SAMPLES_DIR = Path(__file__).parent / "sample_contracts"
SAMPLES_DIR.mkdir(exist_ok=True)

NDA_TEXT = """MUTUAL NON-DISCLOSURE AGREEMENT

This Mutual Non-Disclosure Agreement (the "Agreement") is entered into as of March 1, 2026 (the "Effective Date") by and between Acme Robotics, Inc., a Delaware corporation ("Acme"), and Example Vendor LLC, a California limited liability company ("Vendor"), each a "Party" and collectively the "Parties."

1. TERM
This Agreement shall remain in effect for an initial term of one (1) year from the Effective Date (i.e., through February 28, 2027), after which it renews automatically unless cancelled in advance.

2. TERMINATION
Either Party may terminate this Agreement for convenience upon thirty (30) days' prior written notice to the other Party.

3. CONFIDENTIALITY
All information shared between the Parties, in any form, shall remain confidential indefinitely, with no exceptions.

4. INDEMNIFICATION
Each Party shall indemnify the other against third-party claims arising from such Party's breach of this Agreement, negligence, or willful misconduct, subject to the limitation of liability in Section 5.

5. LIMITATION OF LIABILITY
Vendor's total liability under this Agreement shall not exceed $500, regardless of the nature of the claim.

6. GOVERNING LAW
This Agreement shall be governed by the laws of the State of Delaware, and the Parties consent to the exclusive jurisdiction of the courts located in Delaware.

7. GENERAL
This Agreement constitutes the entire understanding between the Parties with respect to its subject matter.

IN WITNESS WHEREOF, the Parties have executed this Agreement as of the Effective Date.
"""

MSA_TITLE = "MASTER SERVICES AGREEMENT"
MSA_PARAGRAPHS = [
    "This Master Services Agreement (the \"Agreement\") is entered into as of January 15, 2026 (the \"Effective Date\") by and between Northbridge Technologies, Inc., a Delaware corporation (\"Company\"), and Vendor Partner Solutions LLC, a California limited liability company (\"Vendor\").",
    "1. TERM AND RENEWAL",
    "This Agreement commences on the Effective Date and continues for an initial term of one (1) year, renewing automatically thereafter for successive one (1) year terms (next renewal on January 15, 2027) unless either Party provides written notice of non-renewal at least forty-five (45) days prior to the end of the then-current term (i.e., no later than November 30, 2026).",
    "2. TERMINATION",
    "Company may terminate this Agreement at any time without notice or cause. Vendor may not terminate except for Company's material breach.",
    "3. CONFIDENTIALITY",
    "Confidential Information excludes information that is or becomes publicly available through no fault of the receiving Party, was already known prior to disclosure, or is independently developed without use of the disclosing Party's Confidential Information. Obligations survive for three (3) years following termination.",
    "4. INDEMNIFICATION",
    "Vendor shall indemnify, defend, and hold harmless Company from any and all claims, damages, and liabilities of any kind whatsoever, without limitation.",
    "5. LIMITATION OF LIABILITY",
    "Neither Party's aggregate liability under this Agreement shall exceed the total fees paid in the twelve (12) months preceding the claim. Neither Party shall be liable for indirect, incidental, or consequential damages.",
    "6. GOVERNING LAW",
    "This Agreement is governed by the laws of the Cayman Islands, and all disputes must be resolved exclusively in that jurisdiction.",
    "7. PAYMENT TERMS",
    "Invoices are due within thirty (30) days of receipt. Late payments accrue interest at 1.5% per month or the maximum rate permitted by law, whichever is lower.",
    "8. GENERAL",
    "This Agreement constitutes the entire understanding between the Parties with respect to its subject matter.",
]


def write_nda_pdf():
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)
    for line in NDA_TEXT.split("\n"):
        if line.strip():
            pdf.set_x(pdf.l_margin)
            pdf.multi_cell(0, 6, line)
        else:
            pdf.ln(4)
    path = SAMPLES_DIR / "sample_nda.pdf"
    pdf.output(str(path))
    return path


def write_msa_docx():
    doc = Document()
    doc.add_heading(MSA_TITLE, level=1)
    for para in MSA_PARAGRAPHS:
        doc.add_paragraph(para)
    path = SAMPLES_DIR / "sample_msa.docx"
    doc.save(str(path))
    return path


if __name__ == "__main__":
    nda_path = write_nda_pdf()
    msa_path = write_msa_docx()
    print(f"Wrote {nda_path}")
    print(f"Wrote {msa_path}")
